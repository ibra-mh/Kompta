"""Excel workbook builder for the collected and deductible VAT statements."""
from __future__ import annotations

from io import BytesIO
from typing import Any, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .models import ExcelExportRequest, PortfolioExcelRequest


SALES_HEADERS = [
    "Ord", "N° Facture", "Date Facture", "Client", "IF Client", "ICE Client",
    "Montant HT", "Taux TVA", "Montant TVA", "Montant TTC",
]
DEDUCTION_HEADERS = [
    "Ord", "N° Facture", "Désignation", "Fournisseur", "IF", "ICE",
    "Montant HT", "Taux TVA", "Montant TVA", "Montant TTC", "Mode Paiement",
    "Date Facture", "Date Paiement", "Prorata %",
]
NAVY = "1E3A8A"
LIGHT_BLUE = "DBEAFE"
CURRENCY_FORMAT = '#,##0.00 "DH"'


def _style_sheet(
    sheet: Worksheet, headers: Sequence[str], rows: Sequence[Sequence[Any]], total_columns: int
) -> int:
    _style_header(sheet, headers)
    for row_number, values in enumerate(rows, start=2):
        for column, value in enumerate(values, start=1):
            sheet.cell(row_number, column, value)

    total_row = len(rows) + 2
    sheet.cell(total_row, 1, "Totaux").font = Font(bold=True)
    for column in (7, 9, 10):
        letter = get_column_letter(column)
        cell = sheet.cell(total_row, column, f"=SUM({letter}2:{letter}{total_row - 1})")
        cell.font = Font(bold=True)
        cell.number_format = CURRENCY_FORMAT
    for cell in sheet[total_row][:total_columns]:
        cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)

    sheet.sheet_view.showGridLines = True
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(total_columns)}{total_row}"
    for column in range(1, total_columns + 1):
        values = [sheet.cell(row, column).value for row in range(1, total_row + 1)]
        width = max(len(str(value)) if value is not None else 0 for value in values) + 2
        sheet.column_dimensions[get_column_letter(column)].width = min(max(width, 10), 30)
    for row in range(2, total_row):
        for column in (7, 9, 10):
            sheet.cell(row, column).number_format = CURRENCY_FORMAT
        sheet.cell(row, 8).number_format = "0.00"
    return total_row


def build_tva_excel(request: ExcelExportRequest) -> bytes:
    workbook = Workbook()
    sales = workbook.active
    sales.title = "TVA_Collectée_Ventes"
    deductions = workbook.create_sheet("TVA_Déductible_Achats")
    summary = workbook.create_sheet("Synthèse_TVA")

    sales_rows = [[
        line.ord, line.num_facture, line.date_facture, line.client or "",
        line.identifiant_fiscal or "", line.ice or "", line.montant_ht,
        line.taux_tva, line.montant_tva, line.montant_ttc,
    ] for line in request.ventes]
    sales_total_row = _style_sheet(sales, SALES_HEADERS, sales_rows, len(SALES_HEADERS))
    for row in range(2, sales_total_row):
        sales.cell(row, 3).number_format = "dd/mm/yyyy"

    deduction_rows = [[
        line.ord, line.num_facture, line.designation, line.fournisseur or "",
        line.identifiant_fiscal or "", line.ice or "", line.montant_ht,
        line.taux_tva, line.montant_tva, line.montant_ttc,
        line.mode_paiement.value, line.date_facture, line.date_paiement, line.prorata,
    ] for line in request.releve.lines]
    deduction_total_row = _style_sheet(
        deductions, DEDUCTION_HEADERS, deduction_rows, len(DEDUCTION_HEADERS)
    )
    for row in range(2, deduction_total_row):
        deductions.cell(row, 12).number_format = "dd/mm/yyyy"
        deductions.cell(row, 13).number_format = "dd/mm/yyyy"

    for sheet in (sales, deductions, summary):
        sheet.sheet_view.showGridLines = True

    summary["A1"] = "Synthèse TVA"
    summary["A1"].fill = PatternFill("solid", fgColor=NAVY)
    summary["A1"].font = Font(color="FFFFFF", bold=True)
    summary.merge_cells("A1:B1")
    summary_rows = [
        ("Company Name", request.company_name),
        ("Tax Period", request.releve.periode),
        ("Régime", request.regime),
        ("Total TVA Collectée", f"='TVA_Collectée_Ventes'!I{sales_total_row}"),
        ("Total TVA Déductible", f"='TVA_Déductible_Achats'!I{deduction_total_row}"),
        ("TVA Nette", "=B6-B7"),
        ("Statut", '=IF(B8>=0,"TVA Nette à Payer","Crédit de TVA à Reporter")'),
    ]
    for row_number, (label, value) in enumerate(summary_rows, start=3):
        summary.cell(row_number, 1, label).font = Font(bold=True)
        summary.cell(row_number, 2, value)
        if row_number in (6, 7, 8):
            summary.cell(row_number, 2).number_format = "#,##0.00"
    for cell in summary[9]:
        cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        cell.font = Font(bold=True)
    summary.column_dimensions["A"].width = 26
    summary.column_dimensions["B"].width = 28
    summary.freeze_panes = "A3"

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_portfolio_tva_excel(request: PortfolioExcelRequest) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Récapitulatif_Portefeuille"
    purchases = workbook.create_sheet("Tous_Achats_Deductions")
    sales = workbook.create_sheet("Toutes_Ventes")

    sales_headers = [
        "Raison Sociale Client", "Ord", "N° Facture", "Date Facture", "Client",
        "IF Client", "ICE Client", "Montant HT", "Taux TVA", "Montant TVA", "Montant TTC",
    ]
    purchase_headers = [
        "Raison Sociale Client", "Ord", "N° Facture", "Désignation", "Fournisseur", "IF", "ICE",
        "Montant HT", "Taux TVA", "Montant TVA", "Montant TTC", "Mode Paiement",
        "Date Facture", "Date Paiement", "Prorata %",
    ]
    sales_rows = []
    purchase_rows = []
    summary_rows = []
    for client in request.clients:
        collected = sum(line.montant_tva for line in client.ventes)
        deductible = sum(line.montant_tva for line in client.achats)
        summary_rows.append([
            client.company_name, client.if_client, client.ice_client, client.regime,
            collected, deductible, client.credit_anterieur,
        ])
        for line in client.ventes:
            sales_rows.append([
                client.company_name, line.ord, line.num_facture, line.date_facture,
                line.client or "", line.identifiant_fiscal or "", line.ice or "",
                line.montant_ht, line.taux_tva, line.montant_tva, line.montant_ttc,
            ])
        for line in client.achats:
            purchase_rows.append([
                client.company_name, line.ord, line.num_facture, line.designation,
                line.fournisseur, line.identifiant_fiscal, line.ice, line.montant_ht,
                line.taux_tva, line.montant_tva, line.montant_ttc, line.mode_paiement,
                line.date_facture, line.date_paiement, line.prorata,
            ])

    _style_portfolio_sheet(sales, sales_headers, sales_rows, date_columns=(4,))
    _style_portfolio_sheet(purchases, purchase_headers, purchase_rows, date_columns=(13, 14))

    summary_headers = [
        "Raison Sociale", "IF", "ICE", "Régime", "Total TVA Collectée",
        "Total TVA Déductible", "Crédit Antérieur", "TVA Nette à Payer / Crédit à Reporter", "Statut",
    ]
    _style_header(summary, summary_headers)
    for row_number, values in enumerate(summary_rows, start=2):
        for column, value in enumerate(values, start=1):
            summary.cell(row_number, column, value)
        summary.cell(row_number, 8, f"=E{row_number}-F{row_number}-G{row_number}")
        summary.cell(row_number, 9, f'=IF(H{row_number}>=0,"TVA Nette à Payer","Crédit à Reporter")')
    summary_total_row = len(summary_rows) + 2
    summary.cell(summary_total_row, 1, "Totaux").font = Font(bold=True)
    for column in (5, 6, 7, 8):
        letter = get_column_letter(column)
        summary.cell(summary_total_row, column, f"=SUM({letter}2:{letter}{summary_total_row - 1})")
    _finish_sheet(summary, len(summary_headers), summary_total_row)

    for sheet in (summary, purchases, sales):
        sheet.sheet_view.showGridLines = True
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _style_header(sheet: Worksheet, headers: Sequence[str]) -> None:
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(1, column, header)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _finish_sheet(sheet: Worksheet, total_columns: int, total_row: int) -> None:
    for cell in sheet[total_row][:total_columns]:
        cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(total_columns)}{total_row}"
    for column in range(1, total_columns + 1):
        values = [sheet.cell(row, column).value for row in range(1, total_row + 1)]
        width = max(len(str(value)) if value is not None else 0 for value in values) + 2
        sheet.column_dimensions[get_column_letter(column)].width = min(max(width, 10), 36)
    for row in range(2, total_row + 1):
        for column in range(1, total_columns + 1):
            if column in (8, 10, 11) or (sheet.title == "Récapitulatif_Portefeuille" and column in (5, 6, 7, 8)):
                sheet.cell(row, column).number_format = CURRENCY_FORMAT


def _style_portfolio_sheet(
    sheet: Worksheet,
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
    date_columns: Sequence[int] = (),
) -> None:
    _style_header(sheet, headers)
    for row_number, values in enumerate(rows, start=2):
        for column, value in enumerate(values, start=1):
            sheet.cell(row_number, column, value)
    total_row = len(rows) + 2
    sheet.cell(total_row, 1, "Totaux")
    for column in (8, 10, 11):
        letter = get_column_letter(column)
        sheet.cell(total_row, column, f"=SUM({letter}2:{letter}{total_row - 1})")
    for row in range(2, total_row):
        for column in date_columns:
            sheet.cell(row, column).number_format = "dd/mm/yyyy"
    _finish_sheet(sheet, len(headers), total_row)