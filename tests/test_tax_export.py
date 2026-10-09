"""
Unit tests for the SIMPL-TVA / SIMPL-IR export engine.

Run with: pytest -v
"""
from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO
import asyncio

import pytest
import sqlite3
from openpyxl import load_workbook
from pydantic import ValidationError
from fastapi import HTTPException

from core.export_service import export_etat_9421, export_releve_deductions
from core.excel_export import build_portfolio_tva_excel, build_tva_excel
from core.liasse_models import (
    BalanceLine,
    BalanceAmountBasis,
    CgncClass,
    LiasseComputeRequest,
    LiasseMappingCatalog,
    LiasseMappingRule,
    LiasseTable,
)
from core.liasse_service import build_simpl_is_xml, compute_liasse, validate_simpl_is_xml
from core.models import (
    CategorieSalarie,
    DeductionLine,
    Etat9421,
    ExcelExportRequest,
    PortfolioExcelRequest,
    ModePaiement,
    ReleveDeductions,
    SalarieLine,
)
from core.validators import validate_releve_deductions
from core.xml_builders import build_etat_9421_xml, build_releve_deductions_xml
from core.xsd_validator import validate_against_schema
from core.journal_service import JournalEntryPost, JournalRepository, JournalLine
from core.ocr_service import OcrDocumentStore
from core.export_safety import EXPORT_SAFETY_HEADER, EXPORT_SAFETY_LABEL
from core.storage import PROJECT_ROOT, resolve_database_path

SCHEMA_TVA = "schemas/releve_deductions.xsd"
SCHEMA_IR = "schemas/etat_9421.xsd"


class TestStorageConfiguration:
    def test_default_database_location_is_independent_of_working_directory(self, monkeypatch):
        monkeypatch.delenv("KOMPTA_DATABASE_PATH", raising=False)
        assert resolve_database_path() == PROJECT_ROOT / "kompta.sqlite3"

    def test_relative_database_override_is_resolved_from_project_root(self, monkeypatch):
        monkeypatch.setenv("KOMPTA_DATABASE_PATH", "data/local.sqlite3")
        assert resolve_database_path() == PROJECT_ROOT / "data" / "local.sqlite3"


class TestLocalSecurity:
    def test_journal_endpoint_replaces_client_supplied_audit_actor(self, monkeypatch, tmp_path):
        monkeypatch.setenv("KOMPTA_DATABASE_PATH", str(tmp_path / "api-test.sqlite3"))
        import api

        captured = {}

        class Repository:
            def post(self, request):
                captured["request"] = request
                return request

        monkeypatch.setattr(api, "journal_repository", Repository())
        request = JournalEntryPost(
            userId="forged-user", clientId="C001", year=2026, journal="ACHATS",
            date="2026-09-10",
            lines=[JournalLine(compte="6125", debit=100), JournalLine(compte="6125", credit=100)],
        )

        posted = api.post_journal_entry(request)

        assert posted.user_id == "local-owner"
        assert captured["request"].user_id == "local-owner"

    def test_live_server_get_and_export_preflight_allow_exact_origin(self, monkeypatch, tmp_path):
        monkeypatch.setenv("KOMPTA_DATABASE_PATH", str(tmp_path / "cors-test.sqlite3"))
        monkeypatch.setenv("KOMPTA_DEV_CORS_ORIGINS", "http://127.0.0.1:5500")
        import main
        from importlib import reload
        from fastapi.middleware.cors import CORSMiddleware

        app = reload(main).app
        cors = next(item for item in app.user_middleware if item.cls is CORSMiddleware)
        assert cors.kwargs["allow_origins"] == ["null", "http://127.0.0.1:5500"]
        assert cors.kwargs["allow_methods"] == ["GET", "POST", "PUT"]
        assert cors.kwargs["expose_headers"] == ["Content-Disposition", "X-Kompta-Export-Status"]
        assert "*" not in cors.kwargs["allow_origins"]

        async def send_request(method, path, origin, preflight_method=None, request_headers=None):
            messages = []
            headers = [(b"origin", origin.encode())]
            if preflight_method:
                headers.extend([
                    (b"access-control-request-method", preflight_method.encode()),
                    (b"access-control-request-headers", (request_headers or "").encode()),
                ])
            scope = {
                "type": "http", "asgi": {"version": "3.0"},
                "http_version": "1.1", "method": method, "scheme": "http",
                "path": path.split("?", 1)[0], "raw_path": path.split("?", 1)[0].encode(),
                "query_string": path.split("?", 1)[1].encode() if "?" in path else b"",
                "root_path": "", "server": ("127.0.0.1", 8000),
                "client": ("127.0.0.1", 12345), "headers": headers,
            }

            async def receive():
                return {"type": "http.request", "body": b"", "more_body": False}

            async def send(message):
                messages.append(message)

            await app(scope, receive, send)
            response = next(message for message in messages if message["type"] == "http.response.start")
            return response["status"], {key.decode().lower(): value.decode() for key, value in response["headers"]}

        get_status, get_headers = asyncio.run(send_request(
            "GET", "/api/journal-entries?client_id=C001&year=2026", "http://127.0.0.1:5500"
        ))
        preflight_status, preflight_headers = asyncio.run(send_request(
            "OPTIONS", "/export-tva-excel", "http://127.0.0.1:5500", "POST", "content-type"
        ))
        denied_status, _ = asyncio.run(send_request(
            "OPTIONS", "/export-tva-excel", "http://127.0.0.1:5501", "POST", "content-type"
        ))

        assert get_status == 200
        assert get_headers["access-control-allow-origin"] == "http://127.0.0.1:5500"
        assert preflight_status == 200
        assert preflight_headers["access-control-allow-origin"] == "http://127.0.0.1:5500"
        assert "POST" in preflight_headers["access-control-allow-methods"]
        assert "content-type" in preflight_headers["access-control-allow-headers"].lower()
        assert denied_status == 400

    def test_production_cors_default_remains_file_origin_only(self, monkeypatch, tmp_path):
        monkeypatch.setenv("KOMPTA_DATABASE_PATH", str(tmp_path / "production-cors.sqlite3"))
        monkeypatch.delenv("KOMPTA_DEV_CORS_ORIGINS", raising=False)
        import main
        from importlib import reload
        from fastapi.middleware.cors import CORSMiddleware

        app = reload(main).app
        cors = next(item for item in app.user_middleware if item.cls is CORSMiddleware)
        assert cors.kwargs["allow_origins"] == ["null"]

    def test_dev_cors_rejects_wildcard_and_non_loopback_origin(self, monkeypatch):
        import main

        monkeypatch.setenv("KOMPTA_DEV_CORS_ORIGINS", "*")
        with pytest.raises(ValueError, match="exact http://localhost"):
            main.local_development_origins()
        monkeypatch.setenv("KOMPTA_DEV_CORS_ORIGINS", "http://192.168.1.25:5500")
        with pytest.raises(ValueError, match="exact http://localhost"):
            main.local_development_origins()


class TestOcrDocumentStore:
    def test_upload_is_retained_for_review_without_accounting_values(self, tmp_path):
        store = OcrDocumentStore(tmp_path / "ocr.sqlite3")
        document = store.upload("C001", 2026, "facture.pdf", "application/pdf", b"%PDF-1.7 source")

        assert document.status == "needs_review"
        assert store.get(document.document_id).status == "needs_review"

    def test_same_file_in_same_client_returns_duplicate_review_state(self, tmp_path):
        store = OcrDocumentStore(tmp_path / "ocr.sqlite3")
        content = b"%PDF-1.7 source"
        original = store.upload("C001", 2026, "facture.pdf", "application/pdf", content)

        duplicate = store.upload("C001", 2026, "copie.pdf", "application/pdf", content)

        assert duplicate.status == "duplicate_suspected"
        assert duplicate.document_id == original.document_id
        assert duplicate.duplicate_matches == [original.document_id]

    def test_declared_mime_must_match_file_signature(self, tmp_path):
        store = OcrDocumentStore(tmp_path / "ocr.sqlite3")

        with pytest.raises(ValueError, match="type réel"):
            store.upload("C001", 2026, "facture.pdf", "application/pdf", b"not a pdf")

    def test_unsafe_filename_and_empty_file_are_rejected(self, tmp_path):
        store = OcrDocumentStore(tmp_path / "ocr.sqlite3")

        with pytest.raises(ValueError, match="Format"):
            store.upload("C001", 2026, "../facture.pdf", "application/pdf", b"%PDF-1.7")
        with pytest.raises(ValueError, match="vide"):
            store.upload("C001", 2026, "facture.pdf", "application/pdf", b"")

    def test_lookup_is_scoped_to_client_and_exercise(self, tmp_path):
        store = OcrDocumentStore(tmp_path / "ocr.sqlite3")
        document = store.upload("C001", 2026, "facture.pdf", "application/pdf", b"%PDF-1.7 source")

        assert store.get(document.document_id, "C001", 2026) is not None
        assert store.get(document.document_id, "C002", 2026) is None
        assert store.get(document.document_id, "C001", 2025) is None

    def test_additive_schema_migration_preserves_existing_documents(self, tmp_path):
        database = tmp_path / "legacy.sqlite3"
        with sqlite3.connect(database) as db:
            db.execute(
                """CREATE TABLE ocr_documents (
                    id INTEGER PRIMARY KEY AUTOINCREMENT, client_id TEXT NOT NULL,
                    fiscal_year INTEGER NOT NULL, filename TEXT NOT NULL,
                    mime_type TEXT NOT NULL, size INTEGER NOT NULL, checksum TEXT NOT NULL,
                    content BLOB NOT NULL, status TEXT NOT NULL,
                    duplicate_matches TEXT NOT NULL DEFAULT '', processing_error TEXT,
                    uploaded_at TEXT NOT NULL, UNIQUE(client_id, checksum)
                )"""
            )
            db.execute(
                """INSERT INTO ocr_documents
                   (client_id, fiscal_year, filename, mime_type, size, checksum,
                    content, status, uploaded_at)
                   VALUES ('C001', 2026, 'facture.pdf', 'application/pdf', 15,
                           'old-checksum', ?, 'needs_review', '2026-01-01T00:00:00Z')""",
                (b"%PDF-1.7 source",),
            )

        store = OcrDocumentStore(database)

        assert store.get(1).filename == "facture.pdf"
        with sqlite3.connect(database) as db:
            assert "extracted_json" in {row[1] for row in db.execute("PRAGMA table_info(ocr_documents)")}


class TestAppendOnlyJournal:
    def test_posted_entry_date_must_be_iso_formatted(self):
        with pytest.raises(ValidationError, match="AAAA-MM-JJ"):
            JournalEntryPost(
                userId="user-1", clientId="C001", year=2026, journal="ACHATS",
                date="10/09/2026",
                lines=[JournalLine(compte="6125", debit=100), JournalLine(compte="6125", credit=100)],
            )

    def _repository(self, tmp_path):
        repository = JournalRepository(tmp_path / "journal.sqlite3")
        repository._sync_catalog(repository._connect(), [])
        with repository._connect() as db:
            db.executemany(
                "INSERT INTO pcm_accounts(code, label, updated_at) VALUES (?, ?, ?)",
                [("6125", "Charges", "now"), ("4411", "Fournisseurs", "now")],
            )
            db.execute("INSERT INTO auxiliary_accounts(code, label, root_code, client_id, account_type, updated_at) VALUES ('44110001', 'Fournisseur', '4411', 'C001', 'Fournisseur', 'now')")
        return repository

    def _request(self):
        return JournalEntryPost(
            userId="user-1", clientId="C001", year=2026, journal="ACHATS", date="2026-09-10",
            lines=[
                JournalLine(compte="6125", libelle="Achat", debit=100),
                JournalLine(compte="4411", auxiliaire="44110001", libelle="Fournisseur", credit=100),
            ],
        )

    def test_numbers_are_allocated_on_post_and_audit_is_immutable(self, tmp_path):
        repository = self._repository(tmp_path)
        first = repository.post(self._request())
        second = repository.post(self._request())
        assert (first.entry_number, second.entry_number) == (1, 2)
        with pytest.raises(sqlite3.IntegrityError):
            with repository._connect() as db:
                db.execute("UPDATE journal_entries SET label='tampered' WHERE id=?", (first.entry_id,))

    def test_reversal_creates_new_entry_without_editing_original(self, tmp_path):
        repository = self._repository(tmp_path)
        original = repository.post(self._request())
        reversal = repository.reverse(original.entry_id, "user-1")
        assert reversal.entry_id != original.entry_id
        with repository._connect() as db:
            assert db.execute("SELECT COUNT(*) FROM audit_logs WHERE action_type='REVERSAL'").fetchone()[0] == 1

    def test_list_entries_is_scoped_to_client_and_exercise(self, tmp_path):
        repository = self._repository(tmp_path)
        repository.post(self._request())
        other_client = self._request().model_copy(update={
            "client_id": "C002",
            "year": 2025,
            "lines": [
                JournalLine(compte="6125", debit=100),
                JournalLine(compte="6125", credit=100),
            ],
        })
        repository.post(other_client)

        entries = repository.list_entries("C001", 2026)

        assert len(entries) == 1
        assert entries[0]["serverEntryId"] == 1
        assert entries[0]["year"] == 2026

    @staticmethod
    def _catalog_entry(client_id, supplier_code="44110002"):
        catalog = [
            {"code": "6125", "label": "Électricité et eau", "type": "parent"},
            {"code": "4411", "label": "Fournisseurs", "type": "parent"},
            {"code": supplier_code, "label": "Rédal Tétouan", "type": "divisionnaire", "parent": "4411",
             "ice": "000123456789012", "identifiant_fiscal": "12345678"},
        ]
        return JournalEntryPost(
            clientId=client_id, year=2026, journal="ACHATS", date="2026-06-01", accountCatalog=catalog,
            lines=[
                JournalLine(compte="6125", debit=6000),
                JournalLine(compte=supplier_code, auxiliaire=supplier_code, credit=6000),
            ],
        )

    def test_supplier_entry_auto_registers_auxiliary_from_catalog(self, tmp_path):
        repository = JournalRepository(tmp_path / "journal.sqlite3")

        posted = repository.post(self._catalog_entry("C001"))

        assert posted.piece_number == "JA-000001"
        with repository._connect() as db:
            row = db.execute("SELECT root_code, client_id, ice, tax_id, account_type FROM auxiliary_accounts").fetchone()
        assert tuple(row) == ("4411", "C001", "000123456789012", "12345678", "Fournisseur")

    def test_shared_supplier_code_is_registered_per_dossier(self, tmp_path):
        repository = JournalRepository(tmp_path / "journal.sqlite3")

        repository.post(self._catalog_entry("C001"))
        repository.post(self._catalog_entry("C002"))
        repository.post(self._catalog_entry("C001"))

        with repository._connect() as db:
            clients = [row[0] for row in db.execute("SELECT client_id FROM auxiliary_accounts ORDER BY client_id")]
        assert clients == ["C001", "C002"]

    def test_unregistered_auxiliary_is_still_rejected(self, tmp_path):
        repository = self._repository(tmp_path)
        request = self._request().model_copy(update={"lines": [
            JournalLine(compte="6125", debit=100),
            JournalLine(compte="4411", auxiliaire="44119999", credit=100),
        ]})

        with pytest.raises(ValueError, match="Invalid auxiliary account"):
            repository.post(request)

    def test_legacy_auxiliary_table_is_migrated_to_per_dossier_key(self, tmp_path):
        database = tmp_path / "legacy.sqlite3"
        with sqlite3.connect(database) as db:
            db.executescript(
                """CREATE TABLE pcm_accounts (code TEXT PRIMARY KEY, label TEXT NOT NULL, parent TEXT,
                       account_type TEXT NOT NULL DEFAULT 'parent', updated_at TEXT NOT NULL);
                   INSERT INTO pcm_accounts(code, label, updated_at) VALUES ('4411', 'Fournisseurs', 'now');
                   CREATE TABLE auxiliary_accounts (code TEXT PRIMARY KEY, label TEXT NOT NULL,
                       root_code TEXT NOT NULL, client_id TEXT NOT NULL, ice TEXT, tax_id TEXT,
                       account_type TEXT NOT NULL, updated_at TEXT NOT NULL,
                       FOREIGN KEY(root_code) REFERENCES pcm_accounts(code));
                   INSERT INTO auxiliary_accounts(code, label, root_code, client_id, account_type, updated_at)
                       VALUES ('44110001', 'Fournisseur', '4411', 'C001', 'Fournisseur', 'now');"""
            )

        repository = JournalRepository(database)

        with repository._connect() as db:
            pk = [row[1] for row in sorted(db.execute("PRAGMA table_info(auxiliary_accounts)"), key=lambda r: r[5]) if row[5]]
            rows = [tuple(row) for row in db.execute("SELECT code, client_id FROM auxiliary_accounts")]
        assert pk == ["client_id", "code"]
        assert rows == [("44110001", "C001")]

    def test_reversal_rejects_client_or_exercise_mismatch(self, tmp_path):
        repository = self._repository(tmp_path)
        original = repository.post(self._request())

        with pytest.raises(ValueError, match="dossier et cet exercice"):
            repository.reverse(original.entry_id, "local-owner", "C002", 2026)
        with pytest.raises(ValueError, match="dossier et cet exercice"):
            repository.reverse(original.entry_id, "local-owner", "C001", 2025)

        assert repository.list_entries("C001", 2026)[0]["serverEntryId"] == original.entry_id


class TestLiasseMappingSchema:
    def test_catalog_covers_the_three_liasse_tables(self):
        rules = [
            LiasseMappingRule(
                dgi_cell_code="BA-01", label="Immobilisations", table=LiasseTable.BILAN_ACTIF,
                selector={"account_prefixes": ("2",), "cgnc_class": CgncClass.IMMOBILISATIONS},
                amount_basis=BalanceAmountBasis.SOLDE_DEBITEUR, order=1,
            ),
            LiasseMappingRule(
                dgi_cell_code="BP-01", label="Capitaux propres", table=LiasseTable.BILAN_PASSIF,
                selector={"account_prefixes": ("1",), "cgnc_class": CgncClass.CAPITAUX_PROPRES},
                amount_basis=BalanceAmountBasis.SOLDE_CREDITEUR, order=1,
            ),
            LiasseMappingRule(
                dgi_cell_code="CPC-01", label="Charges d'exploitation", table=LiasseTable.CPC,
                selector={"account_prefixes": ("6",), "cgnc_class": CgncClass.CHARGES},
                amount_basis=BalanceAmountBasis.MOUVEMENT_DEBIT, order=1,
            ),
        ]
        catalog = LiasseMappingCatalog(rules=tuple(rules))

        assert len(catalog.rules_for_table(LiasseTable.BILAN_ACTIF)) == 1
        assert catalog.match_account("2355")[0].dgi_cell_code == "BA-01"
        assert catalog.match_account("6111")[0].table == LiasseTable.CPC

    def test_catalog_rejects_duplicate_dgi_cells(self):
        rule = LiasseMappingRule(
            dgi_cell_code="BA-01", label="Immobilisations", table=LiasseTable.BILAN_ACTIF,
            selector={"account_prefixes": ("2",), "cgnc_class": CgncClass.IMMOBILISATIONS},
            amount_basis=BalanceAmountBasis.SOLDE_DEBITEUR, order=1,
        )
        with pytest.raises(ValidationError):
            LiasseMappingCatalog(rules=(rule, rule.model_copy(update={"order": 2})))

    def test_rule_rejects_wrong_table_for_cgnc_class(self):
        with pytest.raises(ValidationError, match="not valid for TABLEAU_1_BILAN_ACTIF; expected one of TABLEAU_3_CPC, TABLEAU_5_ESG"):
            LiasseMappingRule(
                dgi_cell_code="BAD-01", label="Charges", table=LiasseTable.BILAN_ACTIF,
                selector={"account_prefixes": ("6",), "cgnc_class": CgncClass.CHARGES},
                amount_basis=BalanceAmountBasis.MOUVEMENT_DEBIT, order=1,
            )


class TestLiasseCalculation:
    def test_balance_to_fiscal_result_and_tax(self):
        request = LiasseComputeRequest(
            fiscalYear=2026,
            identifiantFiscal="12345678",
            balance=[
                BalanceLine(accountCode="6111", movementDebit=1000),
                BalanceLine(accountCode="7111", movementCredit=3000),
            ],
            adjustments=[
                {"code": "PEN", "label": "Pénalités", "amount": 100, "direction": "reintegrations"},
                {"code": "DIV", "label": "Dividendes", "amount": 200, "direction": "deductions"},
            ],
            mappingRules=[{
                "dgiCellCode": "CPC-01", "label": "Charges", "table": "TABLEAU_3_CPC",
                "selector": {"accountPrefixes": ["6"], "cgncClass": 6},
                "amountBasis": "mouvement_debit", "order": 1,
            }],
        )

        result = compute_liasse(request)

        assert result["tax_adjustments"]["accounting_result"] == Decimal("2000.00")
        assert result["tax_adjustments"]["fiscal_result"] == Decimal("1900.00")
        assert result["tax"]["is"] == Decimal("190.00")
        assert result["tax"]["cotisation_minimale"] == Decimal("15.00")
        assert result["tax"]["tax_due"] == Decimal("190.00")

    def test_simpl_is_xml_contains_required_structural_tags(self):
        request = LiasseComputeRequest(fiscalYear=2026, identifiantFiscal="12345678")
        xml = build_simpl_is_xml(compute_liasse(request))
        assert xml.startswith(b"<?xml")
        assert b"<IdentifiantFiscal>12345678</IdentifiantFiscal>" in xml
        assert b"<ValeursTableau>" in xml
        assert b"<Tableau03>" in xml
        assert b"<Tableau04>" in xml
        assert validate_simpl_is_xml(xml)[0]


def make_valid_line(**overrides) -> DeductionLine:
    defaults = dict(
        ord=1,
        numFacture="FA-2026-001",
        designation="Achat fournitures",
        montantHT=Decimal("1000.00"),
        tauxTVA=Decimal("20"),
        montantTVA=Decimal("200.00"),
        montantTTC=Decimal("1200.00"),
        ice="123456789012345",
        modePaiement=ModePaiement.VIREMENT,
        datePaiement=date(2026, 8, 15),
        dateFacture=date(2026, 8, 10),
    )
    defaults.update(overrides)
    return DeductionLine(**defaults)


# --------------------------------------------------------------------------
# Model-level validation
# --------------------------------------------------------------------------

class TestDeductionLineValidation:
    def test_valid_line_passes(self):
        line = make_valid_line()
        assert line.ord == 1

    def test_ice_must_be_15_digits(self):
        with pytest.raises(ValidationError):
            make_valid_line(ice="12345")  # too short

    def test_if_must_be_8_digits(self):
        with pytest.raises(ValidationError):
            make_valid_line(ice=None, identifiantFiscal="1234")

    def test_missing_ice_and_if_rejected(self):
        with pytest.raises(ValidationError):
            make_valid_line(ice=None, identifiantFiscal=None)

    def test_if_alone_is_sufficient(self):
        line = make_valid_line(ice=None, identifiantFiscal="12345678")
        assert line.identifiant_fiscal == "12345678"

    def test_arithmetic_check_ht_plus_tva_must_equal_ttc(self):
        with pytest.raises(ValidationError):
            make_valid_line(montantHT=Decimal("1000.00"), montantTVA=Decimal("200.00"),
                             montantTTC=Decimal("1300.00"))

    def test_arithmetic_check_tolerates_rounding_to_2dp(self):
        # 1000.005 + 200.005 rounds to 1200.01 -> should equal TTC 1200.01
        line = make_valid_line(
            montantHT=Decimal("1000.005"),
            montantTVA=Decimal("200.005"),
            montantTTC=Decimal("1200.01"),
        )
        assert line.montant_ttc == Decimal("1200.01")


class TestReleveDeductionsValidation:
    def _releve(self, lines):
        return ReleveDeductions(
            ice_declarant="000111222333444",
            if_declarant="12345678",
            periode="2026-08",
            lines=lines,
        )

    def test_valid_releve_no_blocking_issues(self):
        releve = self._releve([make_valid_line(ord=1), make_valid_line(ord=2, numFacture="FA-2026-002")])
        report = validate_releve_deductions(releve)
        assert not report.is_blocked

    def test_duplicate_invoice_detected(self):
        releve = self._releve([
            make_valid_line(ord=1, numFacture="FA-2026-001"),
            make_valid_line(ord=2, numFacture="FA-2026-001"),  # same invoice + same ice
        ])
        report = validate_releve_deductions(releve)
        assert report.is_blocked
        codes = [i.code for i in report.blocking_issues]
        assert "DUPLICATE_INVOICE" in codes

    def test_same_invoice_number_different_ice_not_duplicate(self):
        releve = self._releve([
            make_valid_line(ord=1, numFacture="FA-001", ice="111111111111111"),
            make_valid_line(ord=2, numFacture="FA-001", ice="222222222222222"),
        ])
        report = validate_releve_deductions(releve)
        assert not report.is_blocked

    def test_empty_releve_is_blocking(self):
        releve = self._releve([])
        report = validate_releve_deductions(releve)
        assert report.is_blocked
        assert report.blocking_issues[0].code == "EMPTY_RELEVE"

    def test_unusual_tva_rate_is_warning_not_blocking(self):
        releve = self._releve([make_valid_line(tauxTVA=Decimal("12"))])
        report = validate_releve_deductions(releve)
        assert not report.is_blocked
        assert any(i.code == "UNUSUAL_TVA_RATE" for i in report.warnings)


# --------------------------------------------------------------------------
# XML generation
# --------------------------------------------------------------------------

class TestXmlBuilders:
    def test_releve_xml_contains_expected_tags(self):
        releve = ReleveDeductions(
            ice_declarant="000111222333444",
            if_declarant="12345678",
            periode="2026-08",
            lines=[make_valid_line()],
        )
        xml_bytes = build_releve_deductions_xml(releve)
        xml_str = xml_bytes.decode("utf-8")

        for tag in [
            "<releveDeductions", 'version="V4.0"', "<iceDeclarant>000111222333444</iceDeclarant>",
            "<ifDeclarant>12345678</ifDeclarant>", "<periode>2026-08</periode>",
            "<ord>1</ord>", "<numFacture>FA-2026-001</numFacture>",
            "<montantHT>1000.00</montantHT>", "<tauxTVA>20.00</tauxTVA>",
            "<montantTVA>200.00</montantTVA>", "<montantTTC>1200.00</montantTTC>",
            "<ice>123456789012345</ice>", "<modePaiement>VIREMENT</modePaiement>",
        ]:
            assert tag in xml_str, f"Missing expected fragment: {tag}"

    def test_releve_xml_total_ttc_is_sum_of_lines(self):
        releve = ReleveDeductions(
            ice_declarant="000111222333444",
            if_declarant="12345678",
            periode="2026-08",
            lines=[
                make_valid_line(ord=1, numFacture="A", montantHT=Decimal("1000"),
                                 montantTVA=Decimal("200"), montantTTC=Decimal("1200")),
                make_valid_line(ord=2, numFacture="B", montantHT=Decimal("500"),
                                 montantTVA=Decimal("100"), montantTTC=Decimal("600")),
            ],
        )
        xml_str = build_releve_deductions_xml(releve).decode("utf-8")
        assert "<montantTotalTTC>1800.00</montantTotalTTC>" in xml_str
        assert "<nombreLignes>2</nombreLignes>" in xml_str

    def test_etat_9421_xml_contains_expected_tags(self):
        etat = Etat9421(
            ice_employeur="000111222333444",
            if_employeur="12345678",
            exercice=2026,
            lines=[
                SalarieLine(
                    ord=1, nom="ALAOUI", prenom="Yassine", cin="AB123456",
                    categorie=CategorieSalarie.PERMANENT,
                    date_entree=date(2020, 1, 1),
                    brut_imposable=Decimal("120000"),
                    ir_retenu=Decimal("18000"),
                    net_paye=Decimal("95000"),
                )
            ],
        )
        xml_str = build_etat_9421_xml(etat).decode("utf-8")
        for tag in [
            "<etat9421", 'version="5.1"', "<exercice>2026</exercice>",
            "<nom>ALAOUI</nom>", "<prenom>Yassine</prenom>", "<cin>AB123456</cin>",
            "<categorie>PERMANENT</categorie>", "<brutImposable>120000.00</brutImposable>",
        ]:
            assert tag in xml_str


# --------------------------------------------------------------------------
# XSD structural validation
# --------------------------------------------------------------------------

class TestXsdValidation:
    def test_generated_releve_xml_is_schema_valid(self):
        releve = ReleveDeductions(
            ice_declarant="000111222333444",
            if_declarant="12345678",
            periode="2026-08",
            lines=[make_valid_line()],
        )
        xml_bytes = build_releve_deductions_xml(releve)
        result = validate_against_schema(xml_bytes, SCHEMA_TVA)
        assert result.valid, result.diagnostics

    def test_malformed_xml_reports_line_and_message(self):
        bad_xml = b"<releveDeductions><entete></broken"
        result = validate_against_schema(bad_xml, SCHEMA_TVA)
        assert not result.valid
        assert result.diagnostics
        assert result.diagnostics[0].line >= 1

    def test_generated_etat_9421_xml_is_schema_valid(self):
        etat = Etat9421(
            ice_employeur="000111222333444",
            if_employeur="12345678",
            exercice=2026,
            lines=[
                SalarieLine(
                    ord=1, nom="ALAOUI", prenom="Yassine", cin="AB123456",
                    categorie=CategorieSalarie.STAGIAIRE,
                    brut_imposable=Decimal("30000"),
                    ir_retenu=Decimal("0"),
                    net_paye=Decimal("30000"),
                )
            ],
        )
        xml_bytes = build_etat_9421_xml(etat)
        result = validate_against_schema(xml_bytes, SCHEMA_IR)
        assert result.valid, result.diagnostics

    def test_cached_schema_is_recompiled_when_file_changes(self, tmp_path):
        import os

        schema = tmp_path / "doc.xsd"
        template = ('<xs:schema xmlns:xs="http://www.w3.org/2001/XMLSchema">'
                    '<xs:element name="doc" type="xs:{}"/></xs:schema>')
        schema.write_text(template.format("integer"))
        assert validate_against_schema(b"<doc>12</doc>", schema).valid
        assert not validate_against_schema(b"<doc>abc</doc>", schema).valid

        schema.write_text(template.format("string"))
        stat = schema.stat()
        os.utime(schema, ns=(stat.st_atime_ns, stat.st_mtime_ns + 1_000_000_000))
        assert validate_against_schema(b"<doc>abc</doc>", schema).valid


# --------------------------------------------------------------------------
# Full pipeline (service layer)
# --------------------------------------------------------------------------

class TestExportService:
    def test_portfolio_excel_contains_client_summary_and_consolidated_tabs(self):
        request = PortfolioExcelRequest(clients=[{
            "companyName": "SARL Test",
            "ifClient": "12345678",
            "iceClient": "123456789012345",
            "regime": "Débit",
            "creditAnterieur": "100",
            "ventes": [{
                "ord": 1,
                "numFacture": "V-1",
                "dateFacture": date(2026, 8, 1),
                "client": "Client A",
                "montantHT": "1000",
                "tauxTVA": "20",
                "montantTVA": "200",
                "montantTTC": "1200",
            }],
            "achats": [{
                "companyName": "SARL Test",
                "numFacture": "A-1",
                "designation": "Fournitures",
                "fournisseur": "Fournisseur A",
                "montantHT": "500",
                "tauxTVA": "20",
                "montantTVA": "100",
                "montantTTC": "600",
                "dateFacture": date(2026, 8, 2),
                "datePaiement": date(2026, 8, 10),
            }],
        }])

        workbook = load_workbook(BytesIO(build_portfolio_tva_excel(request)), data_only=False)
        summary = workbook["Récapitulatif_Portefeuille"]
        purchases = workbook["Tous_Achats_Deductions"]
        sales = workbook["Toutes_Ventes"]

        assert workbook.sheetnames == ["Récapitulatif_Portefeuille", "Tous_Achats_Deductions", "Toutes_Ventes"]
        assert summary["A2"].value == "SARL Test"
        assert summary["H2"].value == "=E2-F2-G2"
        assert summary["I2"].value.startswith("=IF(H2")
        assert summary["E2"].number_format == '#,##0.00 "DH"'
        assert purchases["A1"].value == "Raison Sociale Client"
        assert purchases["A2"].value == "SARL Test"
        assert purchases["H3"].value == "=SUM(H2:H2)"
        assert sales["A1"].value == "Raison Sociale Client"
        assert sales["A2"].value == "SARL Test"
        assert sales["H3"].value == "=SUM(H2:H2)"

    def test_tva_excel_contains_tables_totals_and_summary(self):
        request = ExcelExportRequest(
            dossierId="C001",
            companyName="SARL Kompta",
            regime="Encaissement",
            tvaCollectee=Decimal("500"),
            ventes=[{
                "ord": 1,
                "numFacture": "VE-2026-001",
                "dateFacture": date(2026, 8, 12),
                "client": "Client Test",
                "identifiantFiscal": "12345678",
                "ice": "123456789012345",
                "montantHT": Decimal("2500"),
                "tauxTVA": Decimal("20"),
                "montantTVA": Decimal("500"),
                "montantTTC": Decimal("3000"),
            }],
            releve=ReleveDeductions(
                ice_declarant="000111222333444",
                if_declarant="12345678",
                periode="2026-08",
                lines=[make_valid_line(fournisseur="Fournisseur Test")],
            ),
        )

        workbook = load_workbook(BytesIO(build_tva_excel(request)), data_only=False)
        sales = workbook["TVA_Collectée_Ventes"]
        deductions = workbook["TVA_Déductible_Achats"]
        summary = workbook["Synthèse_TVA"]

        assert workbook.sheetnames == ["TVA_Collectée_Ventes", "TVA_Déductible_Achats", "Synthèse_TVA"]
        assert sales["D1"].value == "Client"
        assert sales["D2"].value == "Client Test"
        assert sales["G3"].value == "=SUM(G2:G2)"
        assert sales["I3"].value == "=SUM(I2:I2)"
        assert sales["J3"].value == "=SUM(J2:J2)"
        assert deductions["D1"].value == "Fournisseur"
        assert deductions["D2"].value == "Fournisseur Test"
        assert deductions["G3"].value == "=SUM(G2:G2)"
        assert deductions["I3"].value == "=SUM(I2:I2)"
        assert deductions["J3"].value == "=SUM(J2:J2)"
        assert summary["B6"].value == "='TVA_Collectée_Ventes'!I3"
        assert summary["B7"].value == "='TVA_Déductible_Achats'!I3"
        assert summary["B8"].value == "=B6-B7"
        assert summary["B9"].value.startswith("=IF(B8")
        assert sales.sheet_view.showGridLines
        assert deductions.sheet_view.showGridLines
        assert summary.sheet_view.showGridLines

    def test_empty_client_excel_export_keeps_formatted_tabs(self):
        request = ExcelExportRequest(
            dossierId="C002",
            companyName="Client sans transactions",
            regime="Débit",
            releve=ReleveDeductions(
                ice_declarant="000111222333444",
                if_declarant="12345678",
                periode="2026",
                lines=[],
            ),
        )

        workbook = load_workbook(BytesIO(build_tva_excel(request)), data_only=False)
        assert workbook.sheetnames == ["TVA_Collectée_Ventes", "TVA_Déductible_Achats", "Synthèse_TVA"]
        assert workbook["TVA_Collectée_Ventes"]["A1"].value == "Ord"
        assert workbook["TVA_Déductible_Achats"]["A1"].value == "Ord"
        assert workbook["Synthèse_TVA"]["A1"].value == "Synthèse TVA"

    def test_valid_releve_produces_zip(self):
        releve = ReleveDeductions(
            ice_declarant="000111222333444",
            if_declarant="12345678",
            periode="2026-08",
            lines=[make_valid_line()],
        )
        outcome = export_releve_deductions(releve)
        assert outcome.success
        assert outcome.zip_bytes is not None
        assert outcome.zip_filename == "SIMPL_TVA_2026-08.zip"

    def test_blocking_business_error_prevents_zip(self):
        releve = ReleveDeductions(
            ice_declarant="000111222333444",
            if_declarant="12345678",
            periode="2026-08",
            lines=[
                make_valid_line(ord=1, numFacture="DUP"),
                make_valid_line(ord=2, numFacture="DUP"),
            ],
        )
        outcome = export_releve_deductions(releve)
        assert not outcome.success
        assert outcome.zip_bytes is None
        assert outcome.business_report.is_blocked

    def test_valid_etat_9421_produces_zip(self):
        etat = Etat9421(
            ice_employeur="000111222333444",
            if_employeur="12345678",
            exercice=2026,
            lines=[
                SalarieLine(
                    ord=1, nom="BENNANI", prenom="Sara", cin="CD654321",
                    categorie=CategorieSalarie.PERMANENT,
                    date_entree=date(2019, 3, 1),
                    brut_imposable=Decimal("150000"),
                    ir_retenu=Decimal("25000"),
                    net_paye=Decimal("110000"),
                )
            ],
        )
        outcome = export_etat_9421(etat)
        assert outcome.success
        assert outcome.zip_filename == "ETAT_9421_2026.zip"

    def test_empty_etat_blocked_before_xml_generation(self):
        etat = Etat9421(ice_employeur="000111222333444", if_employeur="12345678", exercice=2026, lines=[])
        outcome = export_etat_9421(etat)
        assert not outcome.success
        assert outcome.business_report.blocking_issues[0].code == "EMPTY_ETAT"


class TestSimplExportSafetyGate:
    def test_unmarked_simpl_exports_are_rejected_before_generation(self):
        import api

        tva = ReleveDeductions(
            ice_declarant="000111222333444", if_declarant="12345678",
            periode="2026-08", lines=[make_valid_line()],
        )
        ir = Etat9421(
            ice_employeur="000111222333444", if_employeur="12345678", exercice=2026,
            lines=[SalarieLine(
                ord=1, nom="BENNANI", prenom="Sara", cin="CD654321",
                categorie=CategorieSalarie.PERMANENT, brut_imposable=Decimal("150000"),
                ir_retenu=Decimal("25000"), net_paye=Decimal("110000"),
            )],
        )
        is_request = LiasseComputeRequest(fiscalYear=2026, identifiantFiscal="12345678")

        for export, payload in (
            (api.export_simpl_tva, tva),
            (api.export_simpl_ir, ir),
            (api.export_simpl_is, is_request),
        ):
            with pytest.raises(HTTPException) as error:
                export(payload)
            assert error.value.status_code == 403
            assert error.value.detail["code"] == "SIMPL_EXPORT_TEST_ONLY"
            assert error.value.detail["message"] == EXPORT_SAFETY_LABEL

    def test_demo_simpl_exports_are_marked_as_test_only(self):
        import api

        tva = ReleveDeductions(
            ice_declarant="000111222333444", if_declarant="12345678",
            periode="2026-08", demoOnly=True, lines=[make_valid_line()],
        )
        ir = Etat9421(
            ice_employeur="000111222333444", if_employeur="12345678", exercice=2026,
            demoOnly=True,
            lines=[SalarieLine(
                ord=1, nom="BENNANI", prenom="Sara", cin="CD654321",
                categorie=CategorieSalarie.PERMANENT, brut_imposable=Decimal("150000"),
                ir_retenu=Decimal("25000"), net_paye=Decimal("110000"),
            )],
        )
        is_request = LiasseComputeRequest(
            fiscalYear=2026, identifiantFiscal="12345678", demoOnly=True,
        )

        for export, payload, expected in (
            (api.export_simpl_tva, tva, "SIMPL_TVA_2026-08.zip"),
            (api.export_simpl_ir, ir, "ETAT_9421_2026.zip"),
            (api.export_simpl_is, is_request, "SIMPL_IS_2026.xml"),
        ):
            response = export(payload)
            assert response.headers["x-kompta-export-status"] == EXPORT_SAFETY_HEADER
            assert f'filename="{expected}"' in response.headers["content-disposition"]

    def test_excel_route_remains_available_without_simpl_demo_marker(self):
        import api

        response = api.export_tva_excel(ExcelExportRequest(
            dossierId="C001", companyName="Client interne", regime="Débit",
            releve=ReleveDeductions(
                ice_declarant="000111222333444", if_declarant="12345678",
                periode="2026-08", lines=[],
            ),
        ))
        assert response.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def test_excel_export_keeps_lines_with_incomplete_supplier_identifiers(self):
        line = {
            "ord": 1, "numFacture": "F1", "designation": "Achat", "montantHT": "100",
            "tauxTVA": "20", "montantTVA": "20", "montantTTC": "120",
            "identifiantFiscal": "IF001", "ice": "001234567", "modePaiement": "VIREMENT",
            "datePaiement": "2025-01-02", "dateFacture": "2025-01-02",
        }
        request = ExcelExportRequest.model_validate({
            "dossierId": "C001", "companyName": "SARL Test", "regime": "Débit",
            "releve": {"periode": "2025", "lines": [line]},
        })

        workbook = load_workbook(BytesIO(build_tva_excel(request)))
        assert workbook["TVA_Déductible_Achats"]["E2"].value == "IF001"
        with pytest.raises(ValidationError):
            ReleveDeductions.model_validate({
                "ice_declarant": "000111222333444", "if_declarant": "12345678",
                "periode": "2025", "lines": [line],
            })
